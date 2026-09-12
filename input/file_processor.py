"""
File Processor Module
Reads uploaded text files and resumes in TXT, PDF, and DOCX formats,
extracts raw text, and performs robust text cleaning and normalization.
"""

import os
import re
from io import BytesIO
from typing import Any, Tuple, Union


def clean_text(raw_text: str) -> str:
    """
    Clean, sanitize, and normalize extracted text:
    - Normalizes unicode quotation marks, hyphens, and whitespace.
    - Strips non-printable ASCII/control characters (keeping tabs and newlines).
    - Collapses repeated redundant whitespace while preserving paragraph flow.
    """
    if not raw_text:
        return ""

    # Replace common unicode typographic characters with standard counterparts
    text = (
        raw_text.replace("\xa0", " ")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\ufeff", "")
    )

    # Remove non-printable control characters (except \n, \r, \t)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "", text)

    # Normalize carriage returns and line feeds
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Collapse horizontal whitespace runs (spaces/tabs) into a single space
    text = re.sub(r"[ \t]+", " ", text)

    # Collapse more than 2 consecutive newlines into 2
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Strip leading and trailing whitespace from each line
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)

    return text.strip()


def extract_from_pdf(file_source: Union[BytesIO, bytes, str, Any]) -> str:
    """
    Extract text content from a PDF file using pypdf or PyPDF2.
    """
    stream = file_source
    if isinstance(file_source, bytes):
        stream = BytesIO(file_source)
    elif hasattr(file_source, "read"):
        if hasattr(file_source, "seek"):
            file_source.seek(0)
        stream = file_source

    try:
        import pypdf
        reader = pypdf.PdfReader(stream)
        pages_text = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                pages_text.append(extracted)
        return "\n\n".join(pages_text)
    except ImportError:
        import PyPDF2
        reader = PyPDF2.PdfReader(stream)
        pages_text = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                pages_text.append(extracted)
        return "\n\n".join(pages_text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}") from e


def extract_from_docx(file_source: Union[BytesIO, bytes, str, Any]) -> str:
    """
    Extract text content from a Microsoft Word DOCX file using python-docx.
    """
    stream = file_source
    if isinstance(file_source, bytes):
        stream = BytesIO(file_source)
    elif hasattr(file_source, "read"):
        if hasattr(file_source, "seek"):
            file_source.seek(0)
        stream = file_source

    try:
        import docx
        doc = docx.Document(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract table text if present
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}") from e


def extract_from_txt(file_source: Union[BytesIO, bytes, str, Any]) -> str:
    """
    Extract text content from plain text file or stream.
    """
    if isinstance(file_source, str):
        if os.path.exists(file_source):
            with open(file_source, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        return file_source

    if hasattr(file_source, "read"):
        if hasattr(file_source, "seek"):
            file_source.seek(0)
        raw_bytes = file_source.read()
    elif isinstance(file_source, bytes):
        raw_bytes = file_source
    else:
        return str(file_source)

    if isinstance(raw_bytes, bytes):
        # Attempt UTF-8 then fallback to Latin-1
        try:
            return raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return raw_bytes.decode("latin-1", errors="replace")

    return str(raw_bytes)


def read_uploaded_file(uploaded_file: Any, filename: str = "") -> str:
    """
    Main function to read and clean text from uploaded files (TXT, PDF, DOCX).

    Args:
        uploaded_file: Streamlit UploadedFile, file-like object, bytes, or file path.
        filename: Optional filename to determine format if not inferrable from uploaded_file.

    Returns:
        str: Cleaned and normalized text.
    """
    fname = filename
    if not fname and hasattr(uploaded_file, "name"):
        fname = uploaded_file.name
    elif not fname and isinstance(uploaded_file, str):
        fname = os.path.basename(uploaded_file)

    ext = os.path.splitext(fname)[1].lower() if fname else ".txt"

    if ext == ".pdf":
        raw = extract_from_pdf(uploaded_file)
    elif ext in [".docx", ".doc"]:
        raw = extract_from_docx(uploaded_file)
    else:
        raw = extract_from_txt(uploaded_file)

    return clean_text(raw)
